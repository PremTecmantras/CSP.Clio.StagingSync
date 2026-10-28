from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = r"C:\Users\yayat\Downloads\SendAnywhere_608146\CLIO (1)\CLIO\CSP.Clio.StagingSync\Clio_Sync_Pipeline_Client_Runbook.docx"

NAVY = "17365D"
BLUE = "D9EAF7"
PALE = "F3F6F9"
GRAY = "D9D9D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def border_table(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = OxmlElement('w:' + edge)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), '4')
        tag.set(qn('w:color'), GRAY)
        borders.append(tag)
    tblPr.append(borders)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        tcPr.append(mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side)
            mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_run_font(run, size=None, bold=None, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def remove_title_rule(paragraph):
    # Word's built-in Title style can inherit a decorative bottom border.
    # Apply explicit "nil" borders on the paragraph so it never renders.
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn('w:pBdr'))
    if existing is not None:
        pPr.remove(existing)
    borders = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right', 'between', 'bar'):
        border = OxmlElement('w:' + edge)
        border.set(qn('w:val'), 'nil')
        borders.append(border)
    pPr.append(borders)

def add_text(paragraph, text, bold=False, size=10.5, color=None):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style='Heading %d' % level)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color="000000")
    return p

def add_body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        add_text(p, item)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    border_table(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        shade(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, label, bold=True, size=9.5, color='FFFFFF')
        if widths: cell.width = Inches(widths[i])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if row_index % 2 == 1: shade(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, value, size=9.5)
            if widths: cell.width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(7)
remove_title_rule(title)
r = title.add_run('Clio Staging Sync and Publishing Runbook')
set_run_font(r, size=24, bold=True, color='000000')

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
add_text(subtitle, 'Client operating guide for the full extraction, validation and reporting publication pipeline', size=12, color='404040')

add_body(doc, 'This guide explains the current approved process for moving Clio data into EKPBI_New_Dev. One EXE run creates a full isolated staging extract, verifies every required endpoint and row count, validates the completed RunID against the current reporting database, and publishes only when every required check passes. The publication refreshes the approved five reporting tables, not every reporting table in the database.')

add_heading(doc, 'What the pipeline does')
add_table(doc,
    ['Stage', 'Purpose', 'Database affected'],
    [
        ('1  Staging extraction', 'Downloads the approved Clio data set into an isolated staging area.', 'Clio_Staging'),
        ('2  Validation', 'Confirms all 19 endpoint results, exact staging counts and the 10 percent comparison control.', 'Read-only check of Clio_Staging and EKPBI_New_Dev'),
        ('3  Publication', 'Copies the approved raw data and refreshes the approved reporting tables.', 'EKPBI_New_Dev'),
    ], [1.45, 3.5, 1.55])

add_heading(doc, 'Normal operating steps')
add_body(doc, 'The normal operation is a single EXE launch with no command-line arguments. It can be started from Visual Studio during testing or from Windows Task Scheduler on the server. The task scheduler must be configured not to start a second instance while a prior run is still active.')

add_table(doc,
    ['Step', 'What the operator does', 'What the system does'],
    [
        ('1', 'Start the scheduled task, or run the application with no command-line arguments.', 'Reads the configured pipeline settings and checks the staging control records.'),
        ('2', 'Wait for completion. Do not start a second copy while the first one is running.', 'Clears the 19 staging raw tables, downloads all pages from Clio and writes the data to Clio_Staging under a unique RunID.'),
        ('3', 'Review the final message and latest run log.', 'Validates the completed RunID. If validation passes, it calls the approved production publishing procedure.'),
        ('4', 'Investigate any failed result before rerunning.', 'Stores endpoint, validation and publication audit information for troubleshooting.'),
    ], [0.45, 2.45, 3.6])

add_heading(doc, 'Automated decision points')
add_body(doc, 'The application does not publish merely because data was downloaded. It must pass each of the following checks before production is changed.')
add_table(doc,
    ['Check', 'Pass condition', 'If it fails'],
    [
        ('Previous run status', 'No previous run remains in RUNNING status, or an earlier stopped run is closed as ABANDONED.', 'The earlier run is recorded as abandoned before the new run begins. Do not schedule overlapping copies of the application.'),
        ('Extraction completion', 'The new RunID is EXTRACTED and every active company and endpoint result is COMPLETE.', 'Production publishing is not started.'),
        ('Received versus saved rows', 'RowsReceived equals RowsWritten for each endpoint/company result.', 'Production publishing is not started.'),
        ('Row-count comparison', 'For each active company and each of the 19 extracted tables, staging is within 10 percent of EKPBI_New_Dev.', 'Production publishing is not started.'),
        ('Duplicate prevention', 'The RunID has not already been published successfully.', 'Production publishing is not started.'),
        ('Publication completion', 'The production audit record finishes with status COMPLETE.', 'The application reports the SQL Server or publisher error for investigation.'),
    ], [1.5, 2.75, 2.25])

add_heading(doc, 'Extracted and validated tables')
add_body(doc, 'The validation follows the same scope as the staging diagnostics process: the 19 tables populated by the Clio staging extractor, compared for each active company. The row-count tolerance is 10 percent. This is a control threshold used to detect unexpected gaps or changes; it is not a replacement for business review of unusual source-data changes.')
add_bullets(doc, [
    'Practice areas, activities, calendars and calendar entries',
    'Relationships, tasks, task types, matters and matter notes',
    'Bills, bank transactions, line items and allocations',
    'Contacts, folders, users, custom fields, custom field sets and activity descriptions',
])

add_heading(doc, 'Staging extraction performance and safety')
add_body(doc, 'Each full run clears the approved 19 staging raw tables before extraction, including CalendarEntries. Clio data is read page by page so large endpoints do not need to remain in memory until the full endpoint download completes. Each completed endpoint is checked against the final committed SQL row count before it can be marked COMPLETE.')
add_body(doc, 'For larger tables, the full extraction uses SQL bulk copy in batches of 1,000 rows. This currently applies to activities, allocations, bank transactions, bills, calendar entries, contacts, line items, matters, matter notes and tasks. The console and the run log show API time, SQL write time, page count, retry information and endpoint throughput to support performance investigation.')

add_heading(doc, 'Publication and reporting refresh')
add_body(doc, 'Only after the new staging RunID passes validation does the application call the approved production stored procedure below. It copies the approved 19 raw tables into EKPBI_New_Dev, records a publication audit row and refreshes the agreed local reporting outputs.')
code = doc.add_paragraph()
code.paragraph_format.left_indent = Inches(0.25)
code.paragraph_format.space_after = Pt(8)
add_text(code, 'dbo.PublishApprovedClioStagingRun  @RunID = new RunID,  @DryRun = 0,  @RebuildReporting = 1', size=9.5, color='17365D')

add_body(doc, 'To avoid a long unnecessary reporting rebuild, the publisher does not call PopulateTables for every active reporting view. The approved allow-list is limited to the five views used by the local CSP.Clio.Sync process:')
add_bullets(doc, [
    'VW_Users',
    'VW_Contacts',
    'VW_Folders',
    'VW_PracticeAreas',
    'VW_Matters',
])
add_body(doc, 'MatterSyncJob-only views and all external API pushes for matters, fees, expenses and documents are outside this pipeline.')

add_heading(doc, 'Success and failure messages')
add_table(doc,
    ['Result', 'Meaning', 'Operator action'],
    [
        ('SYNC VALIDATION PASSED', 'Staging extraction is complete and within the configured validation threshold.', 'No separate manual publication command is required. The application continues to production publication.'),
        ('SYNC VALIDATION FAILED', 'The RunID did not meet one or more extraction or row-count checks.', 'Do not manually run the publication procedure. Review the displayed endpoint, company or count issue.'),
        ('PUBLISH COMPLETE', 'The production audit record and approved reporting refresh completed successfully.', 'Confirm the scheduler result and continue normal reporting monitoring.'),
        ('PUBLISH FAILED', 'SQL Server or the publisher procedure reported an error.', 'Use the displayed SQL error number, procedure and line number to investigate. Do not rerun until the cause is understood.'),
    ], [1.55, 2.65, 2.3])

add_heading(doc, 'Configuration used by the scheduled run')
add_body(doc, 'The application configuration controls the no-argument Visual Studio and Task Scheduler behavior. The following settings are currently enabled:')
add_table(doc,
    ['Setting', 'Current value', 'Meaning'],
    [
        ('StagingDatabaseName', 'Clio_Staging', 'The only database that the extraction EXE is permitted to clear and load.'),
        ('ComparisonDatabaseName', 'EKPBI_New_Dev', 'The database used for validation and publication.'),
        ('AbandonPreviousRunningRun', 'true', 'An earlier RUNNING record is marked ABANDONED before starting a new run. This is intended for a stopped or interrupted earlier run.'),
        ('PublishAfterValidation', 'true', 'A passed new RunID is automatically sent to production publication.'),
    ], [2.15, 1.45, 2.9])
add_body(doc, 'Bulk batch sizes are configured at 1,000 rows for the high-volume full-extraction tables. The setting MaxParallelCompanies exists for future controlled parallel processing, but the current approved run processes companies one at a time. Important: because an earlier RUNNING record is automatically abandoned, the scheduled task must not allow overlapping executions.')

add_heading(doc, 'When a run stops unexpectedly')
add_body(doc, 'If the machine, application or connection stops before the run finishes, the next scheduled launch checks control.SyncRun. With the current setting, it records the unfinished run as ABANDONED and marks unfinished endpoint entries as failed. It then starts a new, independent RunID. It does not publish the abandoned run. Do not publish an ABANDONED or FAILED RunID manually.')

add_heading(doc, 'Run log and audit records')
add_body(doc, 'Every EXE run replaces the previous text log at Logs\\CSP.Clio.StagingSync-last-run.txt beside the EXE. The log mirrors console output, including endpoint timings, validation results and SQL Server publication errors. Database audit evidence is retained in Clio_Staging.control.SyncRun and Clio_Staging.control.SyncEndpointResult, and in EKPBI_New_Dev.control.PublishRun and control.PublishTableAudit.')

add_heading(doc, 'Manual recovery publication')
add_body(doc, 'Manual publication is only for a completed EXTRACTED RunID when the EXE could not reach the publisher, for example because the publisher procedure had not yet been deployed. First run the publisher with @DryRun = 1, review its output, then run it with @DryRun = 0 and @RebuildReporting = 1. Do not manually publish a run that failed validation, is not EXTRACTED, or already has a COMPLETE publication audit record.')

add_heading(doc, 'One time server deployment')
add_bullets(doc, [
    'Deploy the database setup and publisher scripts as approved by the DBA, including the CalendarEntries update to SP_ClearAllData for an existing staging database.',
    'Build the Release configuration in Visual Studio and copy the complete bin\\Release folder to the server.',
    'Update the deployed EXE configuration with the server SQL instance, Clio_Staging, EKPBI_New_Dev and the approved dedicated staging SQL login.',
    'Run one controlled full test and confirm the run log, SyncRun audit, PublishRun audit and five reporting view tables before enabling the scheduler.',
])

add_heading(doc, 'What support needs for investigation')
add_bullets(doc, [
    'The RunID shown in the application or scheduled-task output',
    'The exact validation or SQL Server error message',
    'The endpoint and company name shown beside a failed validation result',
    'The approximate start time and whether the process was interrupted or overlapped with another scheduled run',
])

add_heading(doc, 'Operational summary')
add_body(doc, 'For normal use, start one scheduled application run and wait for its final result. A successful run extracts Clio data to staging, validates every endpoint and row-count control, publishes only when validation passes, and refreshes the five approved local reporting tables. A failed run stops before production publication or records the precise error for investigation in the run log and database audits.')

footer = section.footer
p = footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_text(p, 'Clio Staging Sync and Publishing Runbook', size=8.5, color='666666')

doc.save(OUT)
print(OUT)
