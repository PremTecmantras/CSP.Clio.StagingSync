from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK


OUT = r"C:\Users\yayat\Downloads\SendAnywhere_608146\CLIO (1)\CLIO\CSP.Clio.StagingSync\Incremental_Sync_Process_Guide.docx"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    pPr.append(keep)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def add_code(doc, text):
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cells[i], "F2F6FA")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8.5)
        
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)):
    st = styles[name]
    st.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
    st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.font.size = Pt(size)
    st.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)

if "Code Block" not in styles:
    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code = styles["Code Block"]
code.font.name = "Consolas"
code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
code.font.size = Pt(8.5)
code.paragraph_format.left_indent = Inches(0.18)
code.paragraph_format.right_indent = Inches(0.12)
code.paragraph_format.line_spacing = 1.0
pPr = code._element.get_or_add_pPr()
shd = OxmlElement("w:shd")
shd.set(qn("w:fill"), "F3F5F7")
pPr.append(shd)

footer = section.footer.paragraphs[0]
footer.text = "CSP Clio Staging Sync | Incremental Sync Process Guide"
footer.runs[0].font.size = Pt(8)
footer.runs[0].font.color.rgb = RGBColor(100, 100, 100)
add_page_number(section.footer.add_paragraph())

title = doc.add_paragraph(style="Title")
title.add_run("Clio Incremental Sync Process Guide")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("Full sync, incremental polling, control tables, validation, and publication")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(70, 70, 70)

doc.add_paragraph("Purpose", style="Heading 1")
doc.add_paragraph(
    "This guide explains how CSP.Clio.StagingSync synchronizes Clio data into Clio_Staging and how the resulting data is validated and published to EKPBI_New_Dev. It is intended for developers, database administrators, and operators who configure the two Task Scheduler jobs or investigate a sync result."
)
doc.add_paragraph(
    "The key operating rule is that incremental sync is an update mechanism for staging, while the full sync remains the authoritative reconciliation and publication path. Setting EnableIncrementalSync=true permits incremental execution; it does not change a normal no-argument run into incremental mode."
)

doc.add_paragraph("Executive summary", style="Heading 1")
add_table(doc, ["Area", "Current behavior", "Operational meaning"], [
    ("Full sync", "No command-line argument", "Calls SP_ClearAllData, extracts all configured endpoints, validates, and can publish."),
    ("Incremental sync", "--incremental", "Uses updated_since and changes selected staging rows without clearing or publishing."),
    ("Incremental enablement", "EnableIncrementalSync=true", "Allows --incremental; it does not select the mode by itself."),
    ("Production publication", "Full snapshot only", "The current publisher expects a completed full SyncRun and is not an incremental publisher."),
    ("Deletion handling", "Full sync only", "An updated_since poll cannot prove that an absent record was deleted."),
], [1.35, 2.25, 3.55])

doc.add_paragraph("1 System architecture", style="Heading 1")
doc.add_paragraph(
    "The application has two separate execution paths. The path is selected by the command-line argument, not by the EnableIncrementalSync setting. Both paths use the same executable and configuration file, but they write different control records and have different publication responsibilities."
)
add_table(doc, ["Execution", "Selection", "Reads", "Writes", "Publishes"], [
    ("Full", "No arguments", "All configured Clio endpoints", "Rebuilds Clio_Staging", "Eligible after validation"),
    ("Incremental", "--incremental", "Selected endpoints with updated_since", "Upserts changed rows in Clio_Staging", "No"),
], [1.0, 1.2, 2.0, 2.1, 1.0])

doc.add_paragraph("2 End to end process", style="Heading 1")
doc.add_paragraph("The normal operating sequence is:")
for text in [
    "The full task runs daily. It starts a full SyncRun, clears the staging raw tables, extracts the complete configured dataset, and records endpoint results.",
    "The incremental task runs separately, normally every 15 minutes. It checks that no full run is active, reads the last successful watermark for each company and endpoint, and polls Clio with updated_since.",
    "Each returned record is replaced by CompanyID plus Clio_Internal_ID, or inserted if it does not already exist. The surrounding staging table is not cleared.",
    "If pagination and saving complete successfully, the endpoint result is marked COMPLETE and its watermark advances. If the endpoint fails, its watermark does not advance.",
    "The next full run rebuilds staging from Clio and remains the authoritative operation for deletions and complete reconciliation.",
]:
    add_number(doc, text)

doc.add_paragraph("Execution decision", style="Heading 2")
add_code(doc, "No arguments  ->  full pipeline  ->  ClearAllData  ->  full extraction  ->  validation  ->  optional publication\n--incremental ->  incremental runner ->  updated_since ->  staging upsert ->  watermark ->  no publication")

doc.add_paragraph("3 Incremental polling behavior", style="Heading 1")
doc.add_paragraph(
    "For each company and enabled endpoint, the runner captures a start watermark and an end time. It subtracts the overlap window from the previous watermark before calling Clio. With the current configuration, the overlap is 10 minutes. The overlap intentionally re-reads a small time range so records changed near the previous polling boundary are not missed."
)
add_code(doc, "from_utc = LastSuccessfulToUTC - IncrementalOverlapMinutes\nto_utc   = current UTC time\nClio request: updated_since=from_utc")
doc.add_paragraph(
    "On the first incremental run for a company and endpoint, no watermark exists. The runner therefore starts at the current UTC time minus IncrementalInitialLookbackHours, which is 24 hours in the supplied configuration. This is why a first run can report many changed rows immediately after a full sync. The full sync does not seed the incremental watermark table."
)
doc.add_paragraph("What happens to a returned row", style="Heading 2")
add_bullet(doc, "Existing CompanyID and Clio_Internal_ID: the current raw row is removed and the latest Clio representation is inserted.")
add_bullet(doc, "New CompanyID and Clio_Internal_ID: a new raw row is inserted.")
add_bullet(doc, "No rows returned: the endpoint can still complete successfully and advance its watermark.")
add_bullet(doc, "Record deleted in Clio: updated_since normally does not return it, so the old staging row remains until full reconciliation.")

doc.add_paragraph("4 Control tables", style="Heading 1")
doc.add_paragraph(
    "The SQL deployment script 10_Create_Incremental_Sync_Control.sql creates four tables in the control schema. Together they provide run history, endpoint-level results, per-company watermarks, and a future home for confirmed deletion candidates."
)

doc.add_paragraph("4.1 control.IncrementalSyncRun", style="Heading 2")
doc.add_paragraph("One row represents one invocation of the incremental runner across all active companies and configured endpoints.")
add_table(doc, ["Column", "Meaning"], [
    ("IncrementalRunID", "Unique identifier for the incremental invocation and parent key for endpoint results."),
    ("StartedAtUTC", "UTC time when the run row was created."),
    ("CompletedAtUTC", "UTC time when the run was closed."),
    ("Status", "RUNNING, COMPLETE, FAILED, or ABANDONED."),
    ("ErrorMessage", "Run-level failure detail, when applicable."),
], [1.7, 5.45])

doc.add_paragraph("4.2 control.IncrementalSyncResult", style="Heading 2")
doc.add_paragraph("One row represents one company and endpoint inside an incremental run. This is the primary table for checking which endpoint produced rows or failed.")
add_table(doc, ["Column", "Meaning"], [
    ("IncrementalResultID", "Identity key for the endpoint result."),
    ("IncrementalRunID", "Parent incremental run."),
    ("CompanyID", "Company whose Clio data was polled."),
    ("Endpoint", "For example Activities, Matters, Tasks, or Contacts."),
    ("WatermarkFromUTC", "Effective lower boundary sent to Clio, including overlap."),
    ("WatermarkToUTC", "Upper boundary captured for the successful poll."),
    ("RowsReceived", "Number of Clio records returned by the endpoint."),
    ("Status", "RUNNING, COMPLETE, FAILED, or SKIPPED according to the schema contract."),
    ("ErrorMessage", "Endpoint-level failure detail."),
    ("StartedAtUTC and CompletedAtUTC", "Endpoint timing information."),
], [2.15, 5.0])

doc.add_paragraph("4.3 control.IncrementalWatermark", style="Heading 2")
doc.add_paragraph("One row exists for each CompanyID and Endpoint after its first successful incremental poll. It is the checkpoint used by the next poll.")
add_table(doc, ["Column", "Meaning"], [
    ("CompanyID", "Company scope of the checkpoint."),
    ("Endpoint", "Endpoint scope of the checkpoint."),
    ("LastSuccessfulToUTC", "The last successful upper boundary. This is the next poll's starting reference."),
    ("LastIncrementalRunID", "The run that last advanced the checkpoint."),
    ("UpdatedAtUTC", "When the checkpoint row was last updated."),
], [2.15, 5.0])

doc.add_paragraph("4.4 control.IncrementalDeletionCandidate", style="Heading 2")
doc.add_paragraph("This table is reserved for a later ID-only deletion reconciliation process. The current incremental runner does not automatically populate or confirm deletions. A suspected deletion must be confirmed by an individual Clio GET before production data is removed.")
add_table(doc, ["Column", "Meaning"], [
    ("DeletionCandidateID", "Identity key."),
    ("CompanyID and Endpoint", "Scope of the candidate."),
    ("ClioInternalID", "Clio record believed to be missing."),
    ("FirstSuspectedAtUTC and LastCheckedAtUTC", "Reconciliation timing."),
    ("Status", "SUSPECTED, CONFIRMED_DELETED, NOT_DELETED, or RETRY."),
    ("ConfirmationDetail", "Evidence or error from the individual record check."),
], [2.15, 5.0])

doc.add_paragraph("5 Full sync audit tables", style="Heading 1")
doc.add_paragraph(
    "Incremental control tables are separate from the full-sync audit contract. Full runs are recorded in control.SyncRun and control.SyncEndpointResult. A full run has Mode='FULL'. Incremental runs are recorded in control.IncrementalSyncRun and do not create a full SyncRun."
)
add_table(doc, ["Table", "Use"], [
    ("control.SyncRun", "One full extraction run, including status and Mode='FULL'."),
    ("control.SyncEndpointResult", "Expected, received, written, and completion state for every full company and endpoint."),
    ("control.PublishRun", "Production publication audit in EKPBI_New_Dev, when the full publisher runs."),
    ("control.PublishLease", "Coordinates a full extraction with a production publication so they do not overlap."),
], [2.0, 5.15])

doc.add_paragraph("6 Scheduling two Task Scheduler jobs", style="Heading 1")
doc.add_paragraph("Use two separate tasks that point to the same executable and working directory.")
add_table(doc, ["Task", "Program", "Arguments", "Suggested schedule", "Effect"], [
    ("Clio Full Sync", "CSP.Clio.StagingSync.exe", "None", "Daily, off peak", "Clear, extract, validate, and optionally publish."),
    ("Clio Incremental Sync", "CSP.Clio.StagingSync.exe", "--incremental", "Every 15 minutes", "Poll selected changed endpoints into staging only."),
], [1.25, 1.7, 1.25, 1.35, 1.6])
doc.add_paragraph("For both tasks, set Start in to the folder containing the executable and set the Task Scheduler overlap policy to Do not start a new instance if the task is already running. The database safety checks also reject an incremental run during a full run and reject a full start while an incremental run is active.")

doc.add_paragraph("7 Validation and publication", style="Heading 1")
doc.add_paragraph(
    "The current 10 percent validation is a full-snapshot safety gate. It compares per-company raw row counts in the newly extracted staging snapshot with the corresponding production counts. If the difference exceeds 10 percent, the full run is not eligible for publication. This is intentional: it prevents an unexpectedly incomplete full extraction from replacing production."
)
doc.add_paragraph(
    "Incremental sync can make staging newer than production because it updates staging without publishing. An updated row normally does not change the count; new records increase staging counts, and deletions are not automatically removed by incremental polling. When the next full sync rebuilds staging, the validation result reflects the real difference between the latest Clio snapshot and the current production snapshot."
)
doc.add_paragraph("Why the current full publisher should not be called by incremental sync", style="Heading 2")
add_bullet(doc, "The publisher is designed for a complete full snapshot and an extracted full RunID.")
add_bullet(doc, "Incremental polling covers only selected endpoints.")
add_bullet(doc, "updated_since cannot safely identify all deletions.")
add_bullet(doc, "Direct publication would bypass the full-snapshot validation and could leave production inconsistent.")
doc.add_paragraph("If near-real-time production updates become a requirement, implement a separate incremental publisher with atomic upserts, deletion reconciliation, its own audit records, and coordination with the full publisher. Do not simply reuse the current full publisher with an incremental RunID.")

doc.add_paragraph("8 Operational checks", style="Heading 1")
doc.add_paragraph("Check the latest incremental run:", style="Heading 2")
add_code(doc, "SELECT TOP 10\n    IncrementalRunID, Status, StartedAtUTC, CompletedAtUTC, ErrorMessage\nFROM control.IncrementalSyncRun\nORDER BY StartedAtUTC DESC;")
doc.add_paragraph("Check endpoint rows and watermark movement:", style="Heading 2")
add_code(doc, "SELECT TOP 50\n    IncrementalRunID, CompanyID, Endpoint, RowsReceived, Status,\n    WatermarkFromUTC, WatermarkToUTC, ErrorMessage\nFROM control.IncrementalSyncResult\nORDER BY StartedAtUTC DESC;")
doc.add_paragraph("Check saved checkpoints:", style="Heading 2")
add_code(doc, "SELECT CompanyID, Endpoint, LastSuccessfulToUTC, LastIncrementalRunID, UpdatedAtUTC\nFROM control.IncrementalWatermark\nORDER BY Endpoint, CompanyID;")
doc.add_paragraph("Check whether the latest full run is eligible for validation:", style="Heading 2")
add_code(doc, "SELECT TOP 10 RunID, Mode, Status, StartedAtUTC, CompletedAtUTC, ErrorMessage\nFROM control.SyncRun\nORDER BY StartedAtUTC DESC;")

doc.add_paragraph("9 Troubleshooting guide", style="Heading 1")
add_table(doc, ["Symptom", "Likely explanation", "Check"], [
    ("Incremental reports many rows on first run", "No watermark exists, so the 24-hour initial lookback is used.", "IncrementalSyncResult.WatermarkFromUTC and IncrementalWatermark."),
    ("Incremental runs full extraction", "The task was started without --incremental.", "Console log and command-line arguments."),
    ("SP_ClearAllData appears in the log", "The no-argument full path ran.", "CSP.Clio.StagingSync-last-run.txt."),
    ("Incremental watermark does not move", "The endpoint failed or pagination did not complete.", "IncrementalSyncResult.Status and ErrorMessage."),
    ("Counts differ after incremental", "Staging is newer than production; new records or deletion limitations may be involved.", "Compare CompanyID and Clio_Internal_ID in both databases."),
    ("Full validation exceeds 10 percent", "The latest full staging snapshot differs materially from production.", "Run validation output and investigate before publication."),
], [1.65, 3.0, 2.3])

doc.add_paragraph("10 Log files", style="Heading 1")
doc.add_paragraph("The application writes logs beneath the executable base directory in a Logs folder. With the current logger change:")
add_table(doc, ["Run type", "File"], [
    ("Full", "Logs\\CSP.Clio.StagingSync-last-run.txt"),
    ("Incremental", "Logs\\CSP.Clio.StagingSync-incremental-last-run.txt"),
], [1.5, 5.65])
doc.add_paragraph("Each file is replaced when that mode starts. The control tables remain the durable history for run status, endpoint results, and watermarks.")

doc.add_paragraph("11 Recommended operating policy", style="Heading 1")
for text in [
    "Keep the daily full sync as the only approved production publication path until an incremental publisher is separately designed and tested.",
    "Run incremental polling only with the --incremental argument and EnableIncrementalSync=true.",
    "Do not clear control tables while a run is active or before retaining the relevant audit history.",
    "Investigate failed endpoint results before assuming the watermark is current.",
    "Treat a large first-run row count as an initial lookback result, not automatically as a data corruption event.",
    "Use the full validation result as a review gate; do not bypass it merely to make row counts match.",
]:
    add_bullet(doc, text)

doc.add_paragraph("Appendix A Configuration reference", style="Heading 1")
add_table(doc, ["Setting", "Current value", "Purpose"], [
    ("EnableIncrementalSync", "true", "Allows --incremental to run."),
    ("IncrementalEndpoints", "Activities, CalendarEntries, Tasks, Matters, Bills, Folders, Contacts", "Endpoints polled incrementally."),
    ("IncrementalOverlapMinutes", "10", "Overlap subtracted from the prior watermark."),
    ("IncrementalInitialLookbackHours", "24", "First-poll lookback when no watermark exists."),
    ("PublishAfterValidation", "true", "Allows the full path to invoke the full publisher after validation."),
], [2.15, 2.6, 2.4])

doc.add_paragraph("Appendix B Deployment prerequisite", style="Heading 1")
doc.add_paragraph("Before enabling the incremental Task Scheduler job, run 10_Create_Incremental_Sync_Control.sql once in Clio_Staging and verify that the four control tables exist. Then perform one controlled manual run with --incremental and confirm that the run and every configured endpoint result are COMPLETE before scheduling it.")

doc.core_properties.title = "Clio Incremental Sync Process Guide"
doc.core_properties.subject = "Full and incremental synchronization operations"
doc.core_properties.author = "CSP Clio Staging Sync"
doc.core_properties.keywords = "Clio, incremental sync, full sync, watermark, control tables, Task Scheduler"
doc.save(OUT)
print(OUT)
